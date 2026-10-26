"""Administrator-only leading log and XLSX export. Does not message public groups."""
import io
import os
import sqlite3
import asyncio
import json
import urllib.request
import time
from datetime import datetime
from zoneinfo import ZoneInfo
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from telegram import InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import CommandHandler, CallbackQueryHandler, ConversationHandler, MessageHandler, filters

KST = ZoneInfo("Asia/Seoul")
TRADE_SYMBOL, TRADE_ENTRY, TRADE_T1, TRADE_T2, TRADE_T3, TRADE_RESULT = range(30, 36)
_MARKET_CACHE = {"at": 0, "data": []}


def connection():
    path = os.getenv("DB_PATH", "applications.db").strip()
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    return conn


def setup():
    with connection() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS leading_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                symbol TEXT NOT NULL,
                entry_price REAL NOT NULL CHECK(entry_price > 0),
                target1 REAL NOT NULL,
                target2 REAL NOT NULL,
                target3 REAL NOT NULL,
                result_pct REAL,
                result_type TEXT NOT NULL DEFAULT 'open',
                created_at TEXT NOT NULL,
                closed_at TEXT,
                creator_id INTEGER NOT NULL
            )
        """)
        conn.commit()


def allowed(update, admin_id):
    return update.effective_chat is not None and update.effective_chat.id == admin_id


def menu():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("➕ 리딩 등록", callback_data="lead:add")],
        [InlineKeyboardButton("📝 수익·손절 기록", callback_data="lead:close")],
        [InlineKeyboardButton("📥 리딩 기록 불러오기", callback_data="lead:exports")],
        [InlineKeyboardButton("📊 진행 중인 리딩", callback_data="lead:open")],
    ])


async def start_admin(update, context):
    await update.message.reply_text("📊 투자자 베이스캠프 | 관리자 리딩 관리\n\n리딩 등록, 결과 기록, 엑셀 내보내기를 선택하세요.", reply_markup=menu())



def fetch_upbit_markets():
    if _MARKET_CACHE["data"] and time.monotonic() - _MARKET_CACHE["at"] < 3600:
        return _MARKET_CACHE["data"]
    req = urllib.request.Request("https://api.upbit.com/v1/market/all?is_details=false", headers={"User-Agent": "BasecampBot/1.0"})
    with urllib.request.urlopen(req, timeout=6) as res:
        data = json.load(res)
    markets = [v for v in data if v.get("market", "").startswith("KRW-")]
    _MARKET_CACHE.update({"at": time.monotonic(), "data": markets})
    return markets


async def search_markets(term):
    try:
        markets = await asyncio.to_thread(fetch_upbit_markets)
    except Exception:
        return None
    t = term.casefold().strip()
    return [v for v in markets if t in v.get("market","").casefold() or t in v.get("korean_name","").casefold() or t in v.get("english_name","").casefold()][:10]


async def pick_symbol(update, context):
    q = update.callback_query
    await q.answer()
    form = context.user_data.get("leading_form")
    if not form:
        await q.message.reply_text("입력 시간이 종료되었습니다. 리딩 등록부터 다시 시작해 주세요.")
        return ConversationHandler.END
    try:
        choice = form.get("search_results", [])[int(q.data.rsplit(":", 1)[-1])]
    except (ValueError, IndexError, TypeError):
        await q.message.reply_text("검색이 만료되었습니다. 다시 검색해 주세요.")
        return TRADE_SYMBOL
    await q.edit_message_reply_markup(reply_markup=None)
    symbol = choice["market"].split("-")[1]
    form["symbol"] = symbol
    await q.message.reply_text(f"✅ {symbol} 선택 완료\n진입가(원)를 숫자로 입력하세요. 예: 2000")
    return TRADE_ENTRY

async def begin_add(update, context):
    q = update.callback_query
    await q.answer()
    await q.edit_message_reply_markup(reply_markup=None)
    context.user_data["leading_form"] = {}
    await q.message.reply_text("🔎 종목 검색 (업비트 KRW)\n\n한글 종목명 또는 티커를 입력하세요.\n예: 리플 / XRP / 비트코인\n검색 결과에서 종목을 선택할 수 있습니다.\n취소: /cancel")
    return TRADE_SYMBOL


async def receive_symbol(update, context):
    term = update.message.text.strip()
    if not term or len(term) > 40:
        await update.message.reply_text("종목명 또는 티커를 40자 이내로 입력해 주세요.")
        return TRADE_SYMBOL
    results = await search_markets(term)
    if results is None:
        await update.message.reply_text("⚠️ 업비트 검색 서버 연결에 실패했습니다. 잠시 후 다시 검색해 주세요.")
        return TRADE_SYMBOL
    if not results:
        await update.message.reply_text("검색 결과가 없습니다. 업비트 원화마켓의 한글명 또는 티커로 다시 검색해 주세요.")
        return TRADE_SYMBOL
    context.user_data["leading_form"]["search_results"] = results
    buttons = [[InlineKeyboardButton(f"{x['korean_name']} ({x['market'].split('-')[1]})", callback_data=f"lead:pick:{i}")] for i,x in enumerate(results)]
    await update.message.reply_text("🔎 검색 결과 (업비트 KRW)\n종목을 선택하세요.", reply_markup=InlineKeyboardMarkup(buttons))
    return TRADE_SYMBOL


def parse_positive(t):
    v = float(t.replace(",", "").replace("%", "").strip())
    if not 0 < v < 1e15:
        raise ValueError()
    return v


async def receive_entry(update, context):
    try:
        p = parse_positive(update.message.text)
    except (ValueError, OverflowError):
        await update.message.reply_text("0보다 큰 진입가를 숫자로 입력해 주세요.")
        return TRADE_ENTRY
    context.user_data["leading_form"]["entry_price"] = p
    await update.message.reply_text("1차 목표 수익률(%)을 입력하세요. 예: 5")
    return TRADE_T1


async def goal_step(update, context, key, current, next_state, next_message):
    try:
        v = parse_positive(update.message.text)
        if v > 100000:
            raise ValueError()
        prev = context.user_data["leading_form"].get("target" + str(int(key[-1])-1))
        if prev is not None and v <= prev:
            raise ValueError()
    except (ValueError, OverflowError):
        await update.message.reply_text("이전 목표보다 높은 양수 수익률(%)을 입력해 주세요.")
        return current
    context.user_data["leading_form"][key] = v
    await update.message.reply_text(next_message)
    return next_state


async def goal1(update, context):
    return await goal_step(update, context, "target1", TRADE_T1, TRADE_T2, "2차 목표 수익률(%)을 입력하세요. 예: 10")


async def goal2(update, context):
    return await goal_step(update, context, "target2", TRADE_T2, TRADE_T3, "3차 목표 수익률(%)을 입력하세요. 예: 15")


async def goal3(update, context):
    try:
        v = parse_positive(update.message.text)
        if v <= context.user_data["leading_form"]["target2"] or v > 100000:
            raise ValueError()
    except (ValueError, OverflowError):
        await update.message.reply_text("2차 목표보다 높은 수익률(%)을 입력해 주세요.")
        return TRADE_T3
    d = context.user_data["leading_form"]
    d["target3"] = v
    now = datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S")
    with connection() as conn:
        cursor = conn.execute(
            """INSERT INTO leading_records
            (symbol, entry_price, target1, target2, target3, created_at, creator_id)
            VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (d["symbol"], d["entry_price"], d["target1"], d["target2"], d["target3"], now, update.effective_user.id),
        )
        record_id = cursor.lastrowid
        conn.commit()
    context.user_data.pop("leading_form", None)
    await update.message.reply_text(
        f"✅ 리딩 #{record_id} 등록 완료\n종목: {d['symbol']}\n진입가: {d['entry_price']:,.8g}\n"
        f"목표: +{d['target1']}% / +{d['target2']}% / +{d['target3']}%\n"
        "※ 현재 이 기능은 등록 및 수익·손절 기록 관리용입니다. 실시간 가격 자동감시는 별도 연동이 필요합니다.",
        reply_markup=menu(),
    )
    return ConversationHandler.END


async def record_result(update, context):
    q = update.callback_query
    await q.answer()
    with connection() as conn:
        rows = conn.execute("SELECT id,symbol,entry_price FROM leading_records WHERE result_type='open' ORDER BY id DESC LIMIT 30").fetchall()
    if not rows:
        await q.message.reply_text("현재 결과를 입력할 진행 중인 리딩이 없습니다.", reply_markup=menu())
        return ConversationHandler.END
    buttons = [[InlineKeyboardButton(f"#{r['id']} {r['symbol']} · {r['entry_price']:,.8g}", callback_data=f"lead:choose:{r['id']}")] for r in rows]
    await q.message.reply_text("결과를 기록할 리딩을 선택하세요.", reply_markup=InlineKeyboardMarkup(buttons))
    return TRADE_RESULT


async def result_choose(update, context):
    q = update.callback_query
    await q.answer()
    rid = int(q.data.split(":")[-1])
    with connection() as conn:
        record = conn.execute("SELECT id,symbol FROM leading_records WHERE id=? AND result_type='open'", (rid,)).fetchone()
    if not record:
        await q.message.reply_text("이미 종료됐거나 찾을 수 없는 리딩입니다.")
        return ConversationHandler.END
    context.user_data["leading_close_id"] = rid
    await q.message.reply_text(f"#{rid} {record['symbol']} 최종 실현 수익률(%) 입력\n수익 예: 7.5\n손절 예: -3\n본전: 0\n취소: /cancel")
    return TRADE_RESULT


async def result_input(update, context):
    rid = context.user_data.get("leading_close_id")
    if not rid:
        await update.message.reply_text("종목을 먼저 선택해 주세요.", reply_markup=menu())
        return ConversationHandler.END
    try:
        pct = float(update.message.text.strip().replace("%", "").replace(",", ""))
        if not -100 <= pct <= 100000:
            raise ValueError()
    except (ValueError, OverflowError):
        await update.message.reply_text("수익률을 숫자로 입력하세요. 수익은 양수, 손절은 음수입니다.")
        return TRADE_RESULT
    kind = "profit" if pct > 0 else "loss" if pct < 0 else "breakeven"
    now = datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S")
    with connection() as conn:
        changed = conn.execute(
            """UPDATE leading_records SET result_pct=?, result_type=?, closed_at=?
            WHERE id=? AND result_type='open'""", (pct, kind, now, rid)
        ).rowcount
        conn.commit()
    context.user_data.pop("leading_close_id", None)
    await update.message.reply_text(
        f"{'✅' if changed else '⚠️'} #{rid} {'최종 결과 저장' if changed else '이미 종료된 리딩'}\n최종 수익률: {pct:+.2f}%",
        reply_markup=menu(),
    )
    return ConversationHandler.END


async def cancel(update, context):
    context.user_data.pop("leading_form", None)
    context.user_data.pop("leading_close_id", None)
    await update.message.reply_text("리딩 입력을 취소했습니다.", reply_markup=menu())
    return ConversationHandler.END


async def callbacks(update, context):
    q = update.callback_query
    await q.answer()
    action = q.data
    if action == "lead:exports":
        await q.message.reply_text("📥 리딩 기록 불러오기\n엑셀 다운로드 범위를 선택하세요.", reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton("📈 수익만", callback_data="lead:xlsx:profit"),
             InlineKeyboardButton("📉 손실만", callback_data="lead:xlsx:loss")],
            [InlineKeyboardButton("📋 전체", callback_data="lead:xlsx:all")],
        ]))
    elif action == "lead:open":
        with connection() as conn:
            rows = conn.execute("SELECT id,symbol,entry_price,target1,target2,target3 FROM leading_records WHERE result_type='open' ORDER BY id DESC LIMIT 30").fetchall()
        msg = "📊 진행 중인 리딩 (최근 30건)\n\n" + ("\n".join(f"#{r['id']} {r['symbol']} | 진입 {r['entry_price']:,.8g} | 목표 +{r['target1']}/+{r['target2']}/+{r['target3']}%" for r in rows) if rows else "등록된 리딩이 없습니다.")
        await q.message.reply_text(msg[:4000])
    elif action.startswith("lead:xlsx:"):
        kind = action.rsplit(":",1)[-1]
        if kind not in {"profit", "loss", "all"}:
            return
        sql = "SELECT * FROM leading_records"
        if kind == "profit":
            sql += " WHERE result_type='profit'"
        elif kind == "loss":
            sql += " WHERE result_type='loss'"
        sql += " ORDER BY id DESC"
        with connection() as conn:
            rows = conn.execute(sql).fetchall()
        wb = Workbook()
        ws = wb.active
        ws.title = {"profit":"수익 기록","loss":"손실 기록","all":"전체 리딩"}[kind]
        ws.append(["기록번호", "종목명", "진입가", "수익률(%)", "결과", "1차 목표(%)", "2차 목표(%)", "3차 목표(%)", "등록일시(KST)", "종료일시(KST)"])
        labels = {"profit":"수익", "loss":"손실", "breakeven":"본전", "open":"진행중"}
        for r in rows:
            ws.append([r["id"],r["symbol"],r["entry_price"],r["result_pct"],labels[r["result_type"]],r["target1"],r["target2"],r["target3"],r["created_at"],r["closed_at"]])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183153")
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, width in {"A":12,"B":17,"C":19,"D":16,"E":13,"F":18,"G":18,"H":18,"I":24,"J":24}.items():
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for idx in [3,5,6,7]:
                row[idx].number_format = '0.00"%"'
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        day = datetime.now(KST).strftime("%Y%m%d_%H%M")
        output.name = f"leading_{kind}_{day}.xlsx"
        await q.message.reply_document(document=output, filename=output.name, caption=f"📥 {ws.title}: {len(rows)}건\n※ 미종료 리딩은 전체 파일에만 포함됩니다.")


def install(app, admin_id):
    setup()
    admin_chat = filters.Chat(admin_id)
    app.add_handler(CommandHandler("start", start_admin, filters=admin_chat))
    add = ConversationHandler(
        entry_points=[CallbackQueryHandler(begin_add, pattern=r"^lead:add$")],
        states={
            TRADE_SYMBOL:[CallbackQueryHandler(pick_symbol, pattern=r"^lead:pick:\d+$"), MessageHandler(admin_chat & filters.TEXT & ~filters.COMMAND, receive_symbol)],
            TRADE_ENTRY:[MessageHandler(admin_chat & filters.TEXT & ~filters.COMMAND, receive_entry)],
            TRADE_T1:[MessageHandler(admin_chat & filters.TEXT & ~filters.COMMAND, goal1)],
            TRADE_T2:[MessageHandler(admin_chat & filters.TEXT & ~filters.COMMAND, goal2)],
            TRADE_T3:[MessageHandler(admin_chat & filters.TEXT & ~filters.COMMAND, goal3)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        per_chat=True, per_user=True,
    )
    close = ConversationHandler(
        entry_points=[CallbackQueryHandler(record_result, pattern=r"^lead:close$")],
        states={TRADE_RESULT:[
            CallbackQueryHandler(result_choose, pattern=r"^lead:choose:\d+$"),
            MessageHandler(admin_chat & filters.TEXT & ~filters.COMMAND, result_input),
        ]},
        fallbacks=[CommandHandler("cancel", cancel)],
        per_chat=True, per_user=True,
    )
    app.add_handler(add)
    app.add_handler(close)
    app.add_handler(CallbackQueryHandler(callbacks, pattern=r"^lead:(exports|open|xlsx:(profit|loss|all))$"))
